from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / 'assets'
NAVY = '#111e35'; STONE = '#efede8'; RED = '#b52943'
FONT = 'C:/Windows/Fonts/arial.ttf'; BOLD = 'C:/Windows/Fonts/arialbd.ttf'
def diagram():
    im=Image.new('RGB',(1800,1120),'white');d=ImageDraw.Draw(im)
    def text(x,y,s,size=26,bold=False,color=NAVY): d.text((x,y),s,font=ImageFont.truetype(BOLD if bold else FONT,size),fill=color)
    def box(x,y,w,h,title,body):
        d.rounded_rectangle((x,y,x+w,y+h),radius=16,fill=STONE,outline='#d9d9d9',width=2)
        text(x+22,y+19,title,27,True)
        for i,line in enumerate(body.split('\n')):text(x+22,y+60+i*31,line,23)
    def arrow(points,color=NAVY):
        d.line(points,fill=color,width=4);(x0,y0),(x,y)=points[-2:]
        if x>x0:tri=[(x,y),(x-13,y-8),(x-13,y+8)]
        elif x<x0:tri=[(x,y),(x+13,y-8),(x+13,y+8)]
        elif y>y0:tri=[(x,y),(x-8,y-13),(x+8,y-13)]
        else:tri=[(x,y),(x-8,y+13),(x+8,y+13)]
        d.polygon(tri,fill=color)
    text(35,24,'ONLINE SCAN FLOW',29,True)
    box(35,95,330,170,'Browser UI','Next.js and React\nURL input and results')
    box(475,95,430,170,'FastAPI service','POST /predict\nPydantic validation')
    arrow([(365,178),(475,178)]);text(382,130,'JSON',20)
    box(1070,65,670,125,'Domain check','Bounded A and AAAA queries using dnspython')
    box(1070,220,670,125,'Cached model','TF IDF features and Random Forest prediction')
    box(1070,375,670,125,'URL rules','Keyword combinations and recognized domains')
    arrow([(905,175),(975,175),(975,127),(1070,127)])
    arrow([(975,175),(975,282),(1070,282)])
    arrow([(975,282),(975,437),(1070,437)])
    box(475,390,430,180,'Risk aggregation','Apply score limits and floors\nReturn reasons and availability')
    arrow([(1070,437),(905,437)])
    arrow([(1405,190),(1765,190),(1765,550),(905,550)])
    arrow([(1740,282),(1785,282),(1785,568),(905,568)])
    arrow([(475,482),(198,482),(198,265)]);text(45,508,'Risk score, DNS state and reasons',22)
    d.line((35,620,1765,620),fill='#d9d9d9',width=2);text(35,650,'OFFLINE TRAINING FLOW',29,True)
    box(35,715,370,190,'Training inputs','OpenPhish feed to CSV\nBuilt in benign URLs\nSynthetic phishing URLs')
    box(535,715,480,190,'Training pipeline','Character and word TF IDF\nRandom Forest with 300 trees\nFit on labeled URL strings')
    box(1145,715,595,190,'Saved model artifact','rf_model.joblib\nLoaded and cached by the API\nRestart service after retraining')
    arrow([(405,800),(535,800)]);arrow([(1015,800),(1145,800)])
    arrow([(1470,715),(1470,592),(1030,592),(1030,282),(1070,282)],RED)
    text(1120,604,'Load saved pipeline at runtime',21,color=RED)
    text(35,962,'The scanner performs DNS lookups but does not fetch the destination page.',26)
    text(35,1008,'No database or scheduled feed synchronization is implemented in this project.',25,color='#716f6a')
    im.save(ASSETS/'architecture.png')
diagram()
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.header_distance=Inches(.28);sec.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(8)
normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.12
for name,size in [('Title',29),('Subtitle',12),('Heading 1',21),('Heading 2',13),('Heading 3',11)]:
    doc.styles[name].font.size=Pt(size);doc.styles[name].paragraph_format.keep_with_next=True
for name in ['Heading 1','Heading 2','Heading 3']:doc.styles[name].font.bold=True
header=sec.header.paragraphs[0];header.text='LINKGUARDIAN AI    PROJECT DOCUMENTATION';header.runs[0].font.size=Pt(8);header.runs[0].font.color.rgb=RGBColor(0,0,0)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT;f.add_run('LinkGuardian AI  |  ').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
def p(text,style=None):
    par=doc.add_paragraph(text,style);par.paragraph_format.widow_control=True;return par
def h(text):doc.add_heading(text,2)
def newpage(title):doc.add_page_break();doc.add_heading(title,1)
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for cell,txt in zip(t.rows[0].cells,headers):cell.text=txt
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell,txt in zip(t.add_row().cells,row):cell.text=str(txt)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,cell in enumerate(row.cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'111E35' if i==0 else ('F2F2F2' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcpr.append(borders);margins=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'95');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.05
                for run in par.runs:run.font.size=Pt(9);run.font.bold=i==0;run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    p('');return t
def code(text):
    par=p(text)
    for run in par.runs:run.font.name='Consolas';run.font.size=Pt(8.5)
def bullets(items):
    for item in items:p(item,'List Bullet')
par=doc.add_paragraph();par.add_run().add_picture(str(ASSETS/'linkguardian-logo.png'),width=Inches(.7))
p('LinkGuardian AI project documentation','Title')
p('System overview architecture and technology rationale','Subtitle')
p('Prepared for project review and technical explanation\nProject author  Viraj Gundewar\nDocumentation date  10 October 2026')
h('Project overview')
p('LinkGuardian AI helps a user assess a suspicious website address before opening it. The application combines a trained URL classifier, DNS checks, and explicit URL rules to return a risk estimate with reasons. Its main value is making warning signals understandable at the moment a user decides whether to click a link.')
p('This document describes the frontend, API, model pipeline, risk rules, setup, tests, and limitations. It also explains why the selected technologies suit this design. The product is a decision aid: it does not verify page content, establish the identity of a site owner, or guarantee that a link is safe.')
h('The problem being addressed')
p('Messages can contain unfamiliar domains, misleading login links, or terms that suggest account recovery, bonus claims, or wallet connection requests. Users may have little context for judging the destination. Opening the page to investigate can expose the user to deceptive content. LinkGuardian provides an initial check using the URL string and DNS rather than visiting the page.')
h('Objectives and users')
bullets(['Provide an accessible browser interface for checking public HTTP and HTTPS addresses.','Combine structural model signals with domain resolution and explicit warning rules.','Explain results and disclose unavailable checks instead of presenting unsupported certainty.','Support demonstration, learning, and integration through a documented JSON API.'])
p('The application is suitable for people reviewing links in messages and for developers studying a small machine learning web application. Organizational use requires additional evaluation and operational controls.')
newpage('System architecture')
p('The frontend and backend are separate services. The browser sends JSON to FastAPI; the backend validates the address, performs DNS and model checks, combines the results, and returns JSON. Model training is a separate operation that writes a saved pipeline to disk.')
par=doc.add_paragraph();par.add_run().add_picture(str(ASSETS/'architecture.png'),width=Inches(6.7))
p('Figure 1  Online scanning and offline training. Arrows show data movement; the highlighted arrow connects the saved training artifact to runtime inference.')
table(['Component','Responsibility'],[('Browser interface','Input, loading and errors, risk explanation, session history, and support page.'),('FastAPI service','Validation, health and prediction endpoints, and allowed browser origins.'),('Analysis pipeline','DNS resolution, cached classifier, keyword rules, and score aggregation.'),('Training utilities','Download feed, prepare labeled examples, train, and save the pipeline.')],[1.55,5.15])
p('There is no database, message queue, user account service, or scheduled training worker. DNS resolvers are external runtime dependencies; OpenPhish is used only when a feed download is explicitly requested.')
newpage('How a scan works')
h('User interaction')
p('The user enters a domain or full URL, pastes it, or chooses a sample. The frontend submits one scan at a time and displays an analyzing state. It clears the old result when a scan starts and stores up to five unique recent results in React state after successful scans.')
h('Request validation')
p('The API accepts a URL string between 1 and 4096 characters. It removes surrounding whitespace and adds https:// when the user supplies a domain without a scheme. The hostname is converted to lowercase ASCII using IDNA encoding, and trailing dots are removed. Paths, queries, and fragments remain part of the normalized URL.')
p('Validation rejects unsupported schemes, missing hostnames, whitespace and control characters, embedded credentials, malformed domain labels, invalid ports, single-label hostnames, selected local suffixes, and nonpublic literal IP addresses. These rules make accepted input predictable. They are not a general network isolation mechanism for every possible DNS response.')
h('Domain resolution')
p('dnspython requests an A record with a timeout of 1.5 seconds and a lifetime of 2.5 seconds. If there are no A records, it tries AAAA records. NXDOMAIN produces false; timeouts and inconclusive failures produce null. A valid public IP literal is treated as present without a DNS lookup. Resolution indicates DNS presence, not HTTP reachability or trustworthiness.')
h('Prediction and explanation')
p('The saved classifier estimates the malicious class probability. Rules then adjust the score: recognized domains can reduce false positives, suspicious keyword combinations raise concern, and unavailable checks prevent a clean low-risk result. The API returns the score, category, DNS state, reasons, and model availability.')
h('Response handling')
p('The frontend has a 15 second scan timeout and a 5 second startup health timeout. It checks expected response fields before display. Validation errors show useful feedback, connection failures offer retry, and clipboard denial explains manual paste. Failed connections change the engine status to unavailable; a successful scan refreshes it.')
newpage('Machine learning design')
h('Training data')
p('The training utility downloads the free OpenPhish text feed using Requests with a 15 second timeout and saves URL lines as a CSV file. Training reads the url column using pandas. These URLs and a small set of synthetic phishing examples receive label 1. Generated URLs from a small recognized-domain list and common paths receive label 0.')
p('The current training function uses a limited generated benign set. It does not create a train and test split or calculate precision and recall. The bundled model is a saved artifact; the repository does not contain a complete dataset snapshot or training provenance for it. The training function therefore should not be treated as a complete record of the artifact’s data history.')
h('Feature engineering')
p('A character TF IDF vectorizer extracts character n-grams of length 3 to 5. These capture recurring pieces of a domain, path, or encoded string. A word TF IDF vectorizer extracts individual tokens and token pairs using a pattern that includes single-character tokens. Both use sublinear term frequency. FeatureUnion combines their sparse numeric features.')
p('TF IDF means term frequency multiplied by inverse document frequency. It reduces the weight of common fragments and gives more weight to fragments that distinguish fewer training examples. A URL is represented as a sparse feature vector rather than sent to a generative language model.')
h('Classifier and saved pipeline')
p('RandomForestClassifier uses 300 trees, maximum depth 25, random_state 42, and n_jobs set to minus one for training. The saved artifact also reports 300 trees and maximum depth 25. Scikit-learn Pipeline keeps vectorizers and classifier together so inference applies the learned transformations used in training.')
p('Joblib writes the pipeline to backend/app/rf_model.joblib. The API loads it through an lru_cache with one entry, avoiding disk deserialization for every scan. Startup attempts to load the model but does not download data or retrain. When loading or prediction fails, heuristic analysis remains available and model_available is false.')
h('Why this model fits the project')
p('A Random Forest is practical for a compact supervised classification demonstration: it accepts engineered features, exposes a probability-like output, and has established tooling in scikit-learn. Disadvantages include potentially large artifacts and sparse-feature costs. A logistic regression baseline should be measured before claiming the forest is the best choice. Neither algorithm can reliably establish safety from a URL string alone.')
newpage('Risk scoring and result interpretation')
p('The risk score begins as round(predicted malicious class probability × 100). When the model is unavailable, the initial score is 10 and later rules raise it to at least 40. The following rules apply in order; caps can be followed by floors, so recognized domains still receive caution when DNS cannot be verified.')
table(['Signal','Score rule','Meaning'],[
('Recognized domain without suspicious combination','Cap at 15','Exact recognized hostname or its www form. Arbitrary subdomains are excluded.'),
('Suspicious combination in decoded URL','Floor at 85','Examples include wallet-connect, hack-otp, secure-login, and claim-bonus.'),
('Domain not found in DNS','Floor at 50','Missing DNS does not independently prove malicious intent.'),
('DNS lookup inconclusive','Floor at 40','The domain could not be verified within the available check.'),
('Model unavailable or prediction failed','Floor at 40','The result discloses limited analysis.')],[1.85,1.15,3.7])
h('Risk categories')
table(['Score','Category','Malicious flag'],[('0 to 34','LOW RISK','false'),('35 to 64','CAUTION','false'),('65 to 100','HIGH RISK','true')],[1.35,2.45,2.9])
p('The malicious flag is based on the final score being at least 65. A false value can accompany CAUTION, so clients should show the risk level, availability, and reasons rather than interpreting false as guaranteed safety.')
h('Examples')
p('A recognized domain with working DNS and no suspicious combination receives a score capped at 15. A missing ordinary domain can receive CAUTION at 50 unless another signal raises the score. A URL containing wallet-connect receives at least 85 even when its hostname is recognized. When model and DNS checks are unavailable, the result reports that uncertainty.')
h('What the score does not establish')
p('The value is a rule-adjusted estimate, not a calibrated probability that a site is fraudulent. Both false positives and false negatives are possible. The UI explains the signals and asks the user to verify the sender before sharing passwords or payment details.')
newpage('Technology stack and selection rationale')
p('The rationale explains how each technology supports this implementation. It is an architectural explanation, not a claim that the repository contains a historical comparison of alternatives or measured performance advantages.')
table(['Technology','Role and reason for choosing it'],[
('Next.js 13 and React 18','Pages Router renders scanner and support routes. React state handles input, scan states, and history. Next.js provides routing and production build output.'),
('TypeScript','Interfaces describe results and props, helping catch frontend inconsistencies at build time. External JSON still needs runtime checks.'),
('CSS and Tailwind CSS 3','CSS variables and responsive rules implement the theme. Tailwind and PostCSS support utility processing; most redesigned layouts use explicit CSS classes.'),
('Lucide React and SVG','Lucide supplies interface icons. A custom SVG provides a sharp, reusable logo and favicon at different sizes.'),
('Python and FastAPI','Python connects to the ML libraries. FastAPI exposes JSON endpoints and generated API documentation, separating model logic from the browser.'),
('Pydantic and Uvicorn','Pydantic normalizes and validates payloads. Uvicorn runs the ASGI backend server.'),
('Scikit learn and joblib','Scikit-learn supplies vectorizers, FeatureUnion, Pipeline, and Random Forest. Joblib saves the trained pipeline, separating training from requests.'),
('pandas and Requests','pandas handles training CSV data; Requests downloads the feed with timeout and status checks. Neither is needed for UI rendering.'),
('dnspython','Configurable resolver time limits and distinct DNS outcomes allow missing domains to be separated from inconclusive lookups.'),
('Docker and Compose','Separate Node and Python images package the services. Compose defines local ports and origins for repeatable setup.'),
('unittest and Playwright','Backend tests cover validation and scoring edge cases. Browser checks exercise live requests, navigation, results, and responsive overflow.')],[1.55,5.15])
p('A database is unnecessary for the present scope because there are no accounts and history is in memory. Persistent history, audit trails, and team access would require explicit storage and privacy decisions.')
h('Tradeoffs to explain in a review')
p('The split frontend and backend makes responsibilities clear but requires API and CORS configuration. Python simplifies ML integration but adds a second deployment runtime. The saved pipeline makes inference independent of live feed availability, while manual updates make model freshness an operational responsibility. These are practical tradeoffs rather than universal advantages.')
newpage('API contract and code organization')
h('Endpoints')
table(['Endpoint','Behavior'],[('GET /health','Returns status as ok or degraded and model_available. A degraded response still uses HTTP 200.'),('POST /predict','Accepts a JSON object containing url and returns the normalized address, score, category, flags, and reasons.'),('GET /docs','Generated interactive API documentation.')],[1.6,5.1])
h('Prediction request and response')
code('POST /predict\nContent-Type: application/json\n\n{"url": "google.com"}')
p('Illustrative low-risk response shape. The score can vary with the model and DNS conditions.')
code('{\n  "url": "https://google.com",\n  "is_malicious": false,\n  "domain_exists": true,\n  "risk_score": 15,\n  "risk_level": "LOW RISK",\n  "reason": "The hostname matches a recognized domain; this does not verify the page content.",\n  "model_available": true\n}')
p('domain_exists can be true, false, or null. Multiple explanations in reason are separated by a vertical bar with surrounding spaces. Invalid payloads return HTTP 422. Unexpected analysis failures return HTTP 503 with a generic message while the server logs the exception. Authentication and rate limiting are not implemented.')
h('Files to understand first')
table(['Path','Purpose'],[
('frontend/pages/index.tsx','Scanner, API requests, results, and recent checks.'),('frontend/pages/contact.tsx','Support page and mailto draft form.'),('frontend/components and styles','Shared navigation, footer, logo, and theme.'),('backend/app/main.py','API app, validation, CORS, startup, and endpoints.'),('backend/app/model.py','Training, inference, DNS, URL rules, and scoring.'),('backend/app/utils.py','OpenPhish download and CSV creation.'),('backend/test_backend.py','Nine backend regression tests.'),('frontend/browser-check.cjs','Browser smoke checks using Microsoft Edge.')],[2.5,4.2])
newpage('User experience branding and data handling')
h('Interface design')
p('The scanner is the primary task, so the address field and Analyze link button sit together at the top of the workspace. Empty, loading, error, low-risk, and caution states each explain what to do next. Results show the estimate, category, DNS state, model availability, and reasons. Recent results can be reopened or cleared without a new scan.')
p('Layouts adapt from desktop columns to a mobile column. Visible labels, keyboard focus indicators, a skip link, live result announcements, descriptive controls, and reduced-motion styling support accessibility. The checks performed so far are not a full accessibility certification.')
h('Logo and colour system')
p('The new mark combines linked shapes inside a guardian outline. The linked forms represent a website address and connection; the enclosing outline represents a protective check. It uses midnight navy #111E35, warm stone grey #EFEDE8, and crimson red #B52943. A standalone SVG is reused in the header, footer, hero graphic, and favicon. Warning colours in results remain separate from the logo’s decorative accent.')
h('Privacy and external communication')
p('Scan history lives in memory in the current browser tab and disappears on reload. The application has no database that stores it. A submitted URL is sent to the configured backend; DNS checks share the hostname with a resolver. The backend analyzes the URL string but does not fetch the destination. This does not imply that hosting infrastructure never logs requests or that the service is anonymous.')
p('The support form builds a mailto URL from the entered details. It prepares an email draft for the user to review and send, rather than sending through the backend. Clipboard access is requested when Paste or Copy is chosen. Remove sensitive tokens, personal information, and credentials from URLs shared as support examples.')
h('Security measures and boundaries')
p('CORS defaults to local frontend origins and allows GET, POST, and Content-Type without credentials. Deployed origins should be explicitly configured. Input validation rejects unsupported inputs, but CORS is not authentication and does not stop direct API clients. Joblib deserializes the model, so artifacts must come from a trusted training process; untrusted files must not be loaded.')
newpage('Setup deployment and operations')
h('Local setup')
p('Use a Python environment that supports backend/requirements.txt and Node.js for the frontend. Containers use Python 3.12 and Node 20. Run these commands from the project root in separate terminals.')
code('python -m pip install -r backend/requirements.txt\npython -m uvicorn backend.app.main:app --host 127.0.0.1 --port 8000')
code('npm --prefix frontend ci\nnpm run dev')
p('The standard UI address is http://localhost:3000 and API documentation is at http://localhost:8000/docs. Root scripts delegate frontend commands into its folder. npm run build creates production assets; npm start serves them.')
h('Environment configuration')
p('NEXT_PUBLIC_API_URL sets the backend base URL and defaults to http://localhost:8000. Configure it in frontend/.env.local. Next.js embeds it at build time, so rebuild after changes. The frontend Dockerfile accepts it as a build argument. A browser must be able to reach that address; a Compose service name alone is not a public browser URL.')
p('CORS_ORIGINS is a comma-separated backend variable. It defaults to http://localhost:3000 and http://127.0.0.1:3000. Set the exact deployed frontend origin, including scheme and port. For remote hosting, configure HTTPS and operational controls at the hosting or proxy layer.')
h('Docker and model refresh')
code('docker compose up --build')
p('Compose exposes UI port 3000 and API port 8000. The backend image includes the model file. The frontend image includes production assets, runtime dependencies, and public logo files. Adjust local defaults for remote hosting.')
code('python -c "from backend.app.utils import download_openphish_dataset; from backend.app.model import train_model; train_model(download_openphish_dataset())"')
p('Run the retraining command on one line in a controlled environment, then restart the backend. Training clears its own in-process cache, but another running service retains its cached object. No timer or automatic synchronization is implemented. Retain the data, version the artifact, and evaluate a replacement before deployment.')
h('Operational checks')
p('Check /health after startup, inspect logs when analysis fails, and verify browser access to the API origin. DNS failures should remain inconclusive rather than proof of fraud. Protect model artifacts and preserve dependency compatibility across environments.')
newpage('Verification limitations and future work')
h('Verification completed')
p('The revised frontend passed its Next.js production build. Nine backend regression tests passed: normalization, invalid inputs, recognized-domain boundaries, missing-model behavior, missing and inconclusive DNS, keyword handling, API validation, and generic error responses. Live browser checks passed for engine status, recognized and suspicious scans, validation, history, support navigation, FAQ, responsive overflow, and absence of JavaScript runtime errors.')
p('Backend tests mock some DNS and model failures for deterministic checks. Live scans demonstrate integration, not predictive accuracy. Compose configuration validation passed; container images have not been built or exercised as part of those checks.')
code('npm run test:backend\nnpm run build\nnpm run test:browser')
p('The browser script defaults to http://localhost:3001 and uses installed Microsoft Edge. Start the production UI on that port and allow its origin through backend CORS_ORIGINS. BROWSER_TEST_URL can select another preview. These are local preview choices, not fixed deployment requirements.')
h('Current limitations')
bullets(['No measured precision, recall, F1 score, confusion matrix, or calibrated probability is provided.','The limited benign training set can bias predictions toward familiar domains. Dataset lineage for the saved artifact is incomplete.','The scanner does not inspect HTML, redirects, TLS certificates, scripts, or downloads. DNS cannot establish legitimacy.','Recognized domains can host compromised content; keyword rules can still produce false positives.','Authentication, rate limiting, accounts, persistent history, automatic feed updates, and server-side email are not implemented.'])
h('Recommended next steps')
p('Establish representative labeled data, deduplicate related domains, and reserve independent evaluation data. Compare a simpler classifier with the forest, measure false positives and negatives, and calibrate scores if probabilities are needed. Then add model versions, reproducible training records, monitoring, and bounded capacity. Persistent history, account access, and page inspection should be separate features with explicit privacy and security requirements.')
h('Basis for this document')
p('Facts were checked against README.md, frontend/package.json, frontend pages and components, backend/requirements.txt, backend/app/main.py, backend/app/model.py, backend/app/utils.py, tests, and Docker configuration. The saved model was inspected for classifier settings. Technology rationale follows these implementation roles; no original decision log or performance comparison is claimed.')
doc.core_properties.title='LinkGuardian AI project documentation'
doc.core_properties.subject='Architecture implementation and technology rationale'
doc.core_properties.author='Viraj Gundewar'
doc.core_properties.keywords='LinkGuardian, URL risk, FastAPI, Next.js, Random Forest'
for shape in doc.inline_shapes:
    shape._inline.docPr.set('descr','LinkGuardian brand mark' if shape.width < Inches(1) else 'Architecture showing browser to API analysis and separate feed to training to saved model flow')
for node in list(doc.styles['Title'].element.xpath('.//w:pBdr')):
    node.getparent().remove(node)
out=ROOT/'LinkGuardian_Project_Documentation.docx';doc.save(out);print(out)
